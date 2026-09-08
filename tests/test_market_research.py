from __future__ import annotations

import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from twextract import excel_export
from twextract.excel_export import write_market_research_excel
from twextract.market_research import (
    drop_duplicate_records,
    nutriscore_sentiment,
    publisher_from_url,
    record_from_hn,
    record_from_openfoodfacts,
    record_identity_keys,
    sentiment_from_engagement,
)


class MarketResearchParseTests(unittest.TestCase):
    def test_sentiment_bands(self) -> None:
        self.assertEqual(sentiment_from_engagement(250, 40), "strong_positive")
        self.assertEqual(sentiment_from_engagement(30), "positive")
        self.assertEqual(sentiment_from_engagement(8), "neutral")
        self.assertEqual(sentiment_from_engagement(1), "low")
        self.assertEqual(sentiment_from_engagement(None), "unscored")
        self.assertEqual(nutriscore_sentiment("a"), "positive")
        self.assertEqual(nutriscore_sentiment("c"), "mixed")
        self.assertEqual(nutriscore_sentiment("e"), "negative")

    def test_publisher_from_url(self) -> None:
        self.assertEqual(publisher_from_url("https://www.example.com/retail"), "example.com")
        self.assertEqual(publisher_from_url("https://news.ycombinator.com/item?id=1"), "Y Combinator")

    def test_openfoodfacts_uses_real_brand(self) -> None:
        row = record_from_openfoodfacts(
            {
                "code": "3017620422003",
                "product_name": "Nutella",
                "brands": "Ferrero, Nutella",
                "categories_tags": ["en:breakfasts", "en:sweet-spreads"],
                "nutriscore_grade": "e",
                "ecoscore_grade": "d",
                "countries_tags": ["en:france", "en:germany"],
                "unique_scans_n": 12000,
                "last_modified_t": 1725148800,
            }
        )
        assert row is not None
        self.assertEqual(row["source"], "world.openfoodfacts.org")
        self.assertEqual(row["catalog"], "Open Food Facts")
        self.assertEqual(row["brand"], "Ferrero")
        self.assertEqual(row["title"], "Nutella")
        self.assertEqual(row["sentiment"], "negative")
        self.assertIn("Nutri-Score E", row["body"])
        self.assertNotIn("dummyjson", str(row).lower())
        self.assertNotIn(None, row.values())

    def test_hn_story_and_comment(self) -> None:
        story = record_from_hn(
            {
                "objectID": "123",
                "title": "Market share in EU retail",
                "points": 140,
                "num_comments": 22,
                "author": "alice",
                "created_at": "2026-09-04T00:00:00.000Z",
                "url": "https://example.com/retail",
                "_tags": ["story"],
            },
            record_type="trend",
            topic="retail",
        )
        comment = record_from_hn(
            {
                "objectID": "456",
                "story_title": "Market share in EU retail",
                "comment_text": "<p>Prices feel high</p>",
                "author": "bob",
                "_tags": ["comment"],
                "parent_id": 123,
            },
            record_type="sentiment",
            topic="hn-comments",
        )
        assert story is not None and comment is not None
        self.assertEqual(story["record_id"], "hn-story-123")
        self.assertEqual(story["brand"], "example.com")
        self.assertEqual(story["source"], "hn.algolia.com")
        self.assertEqual(story["sentiment"], "strong_positive")
        self.assertEqual(comment["record_id"], "hn-comment-456")
        self.assertEqual(comment["title"], "Market share in EU retail")
        self.assertEqual(comment["body"], "Prices feel high")
        self.assertEqual(comment["author"], "bob")
        self.assertEqual(comment["brand"], "Y Combinator")
        self.assertEqual(comment["score"], 0)
        self.assertEqual(comment["comment_count"], 0)
        self.assertEqual(comment["published_at"], "not provided by source")
        self.assertNotIn("dummyjson", str(story).lower())
        self.assertNotIn(None, story.values())
        self.assertNotIn(None, comment.values())
        self.assertFalse(record_identity_keys(story) & record_identity_keys(comment))

    def test_drops_duplicate_ids(self) -> None:
        first = record_from_hn(
            {"objectID": "1", "title": "A", "points": 1, "_tags": ["story"]},
            record_type="trend",
            topic="hn-stories",
        )
        second = record_from_hn(
            {"objectID": "1", "title": "A copy", "points": 9, "_tags": ["story"]},
            record_type="trend",
            topic="retail",
        )
        unique = drop_duplicate_records([first, second])
        self.assertEqual(len(unique), 1)


class MarketResearchExcelTests(unittest.TestCase):
    def test_write_workbook(self) -> None:
        rows = [
            {
                "row_number": 1,
                "ok": True,
                "use_case": "Market Research",
                "record_type": "trend",
                "topic": "retail",
                "title": "EU grocery prices",
                "sentiment": "positive",
                "source": "hn.algolia.com",
                "brand": "example.com",
                "exit": {"ip": "8.8.8.8", "city": "Dallas"},
            }
        ]
        summary = [{"metric": "Records scraped", "value": 1}]
        with TemporaryDirectory() as tmp:
            folder = Path(tmp)
            excel_export.OUTPUT_DIR = folder
            excel_export.MARKET_RESEARCH_XLSX = folder / "market_research.xlsx"
            path = write_market_research_excel(rows, summary)
            from openpyxl import load_workbook

            book = load_workbook(path)
            self.assertEqual(path.name, "market_research.xlsx")
            self.assertIn("Market Research", book.sheetnames)
            self.assertIn("Research summary", book.sheetnames)
            headers = [cell.value for cell in book["Market Research"][1]]
            self.assertIn("title", headers)
            self.assertIn("sentiment", headers)
            self.assertIn("record_type", headers)
            self.assertEqual(headers[-1], "complete Json")


if __name__ == "__main__":
    unittest.main()
