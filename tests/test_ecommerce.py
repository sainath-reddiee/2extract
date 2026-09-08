from __future__ import annotations

import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from twextract import excel_export
from twextract.ecommerce import (
    interleave_buckets,
    parse_availability,
    parse_category_list,
    parse_price,
    parse_rating,
    product_from_dummyjson,
    product_from_openprices,
)
from twextract.excel_export import flatten_ecommerce_rows, write_ecommerce_excel


DUMMYJSON_PRODUCT = {
    "id": 121,
    "title": "iPhone 5s",
    "description": "A classic smartphone.",
    "category": "smartphones",
    "price": 199.99,
    "discountPercentage": 12.91,
    "rating": 2.83,
    "stock": 25,
    "tags": ["smartphones", "apple"],
    "brand": "Apple",
    "sku": "SMA-APP-IPH-121",
    "availabilityStatus": "In Stock",
    "warrantyInformation": "Lifetime warranty",
    "shippingInformation": "Ships in 1 month",
    "returnPolicy": "60 days return policy",
    "thumbnail": "https://cdn.dummyjson.com/phone.webp",
    "meta": {"barcode": "8814683940853"},
}

OPEN_PRICE_ITEM = {
    "product_code": "3017620422003",
    "product_name": "NUTELLA",
    "price": 3.49,
    "currency": "EUR",
    "date": "2026-09-03",
    "price_is_discounted": False,
    "product": {
        "code": "3017620422003",
        "product_name": "Nutella",
        "brands": "Ferrero",
        "quantity": "400 g",
        "categories_tags": ["en:breakfasts", "en:spreads", "en:sweet-spreads"],
        "image_url": "https://images.openfoodfacts.org/nutella.jpg",
    },
    "location": {
        "osm_brand": "Carrefour",
        "osm_name": "Carrefour Paris",
        "osm_address_city": "Paris",
        "osm_address_country_code": "FR",
    },
}


class EcommerceParseTests(unittest.TestCase):
    def test_parse_price(self) -> None:
        raw, value, currency = parse_price("£51.77")
        self.assertEqual(raw, "£51.77")
        self.assertEqual(value, 51.77)
        self.assertEqual(currency, "GBP")
        _raw, dollar, usd = parse_price("$199.99")
        self.assertEqual(dollar, 199.99)
        self.assertEqual(usd, "USD")

    def test_parse_availability(self) -> None:
        raw, in_stock, count = parse_availability("In stock (22 available)")
        self.assertTrue(in_stock)
        self.assertEqual(count, 22)
        self.assertIn("In stock", raw or "")
        _raw, low, _count = parse_availability("Low Stock")
        self.assertTrue(low)

    def test_parse_rating(self) -> None:
        label, value = parse_rating(["star-rating", "Three"])
        self.assertEqual(label, "Three")
        self.assertEqual(value, 3)

    def test_parse_category_list(self) -> None:
        parsed = parse_category_list(
            [
                {"slug": "smartphones", "name": "Smartphones", "url": "https://dummyjson.com/products/category/smartphones"},
                "laptops",
                {"slug": "smartphones", "name": "dup"},
            ]
        )
        self.assertEqual([row["slug"] for row in parsed], ["smartphones", "laptops"])

    def test_product_from_dummyjson(self) -> None:
        row = product_from_dummyjson(DUMMYJSON_PRODUCT)
        self.assertEqual(row["product_name"], "iPhone 5s")
        self.assertEqual(row["category"], "smartphones")
        self.assertEqual(row["price_value"], 199.99)
        self.assertEqual(row["currency"], "USD")
        self.assertTrue(row["in_stock"])
        self.assertEqual(row["stock_count"], 25)
        self.assertEqual(row["upc"], "8814683940853")
        self.assertEqual(row["source"], "dummyjson.com")

    def test_product_from_openprices(self) -> None:
        row = product_from_openprices(OPEN_PRICE_ITEM, listing_page=1)
        assert row is not None
        self.assertEqual(row["product_name"], "NUTELLA")
        self.assertEqual(row["price_value"], 3.49)
        self.assertEqual(row["currency"], "EUR")
        self.assertEqual(row["price"], "€3.49")
        self.assertEqual(row["brand"], "Ferrero")
        self.assertEqual(row["stores"], "Carrefour")
        self.assertTrue(row["has_price"])
        self.assertTrue(row["in_stock"])
        self.assertEqual(row["availability"], "In stock at Carrefour, Paris, FR on 2026-09-03")
        self.assertIsNone(product_from_openprices({"product_name": "X"}))
        self.assertIsNone(product_from_openprices({"price": 1.2, "currency": "EUR"}))

    def test_enrich_openprices_availability(self) -> None:
        from twextract.ecommerce import enrich_openprices_availability

        row = {
            "source": "prices.openfoodfacts.org",
            "stores": "Netto Marken-Discount",
            "store_city": "Berlin",
            "store_country": "DE",
            "price_date": "2026-09-04",
            "availability": None,
            "in_stock": None,
        }
        enrich_openprices_availability(row)
        self.assertTrue(row["in_stock"])
        self.assertIn("Netto Marken-Discount", row["availability"])
        self.assertIn("Berlin", row["availability"])

    def test_interleave_buckets_mixes_categories(self) -> None:
        phones = [
            product_from_dummyjson(
                {**DUMMYJSON_PRODUCT, "id": 1, "title": "Phone A", "category": "smartphones", "sku": "PHN-1", "meta": {"barcode": "111"}}
            )
        ]
        laptops = [
            product_from_dummyjson(
                {**DUMMYJSON_PRODUCT, "id": 2, "title": "Laptop A", "category": "laptops", "sku": "LAP-1", "meta": {"barcode": "222"}}
            )
        ]
        mixed = interleave_buckets([phones, laptops], 2)
        self.assertEqual([row["category"] for row in mixed], ["smartphones", "laptops"])

    def test_skips_duplicate_name_and_brand_across_sources(self) -> None:
        from twextract.ecommerce import drop_duplicate_products, is_duplicate_product, product_identity_keys

        dummy = product_from_dummyjson(
            {**DUMMYJSON_PRODUCT, "title": "Nutella", "brand": "Ferrero", "meta": {"barcode": "111"}}
        )
        price = product_from_openprices(OPEN_PRICE_ITEM)
        assert price is not None
        price["product_name"] = "NUTELLA"
        price["brand"] = "Ferrero"
        price["upc"] = "999"
        self.assertTrue(product_identity_keys(dummy) & product_identity_keys(price))
        unique = drop_duplicate_products([dummy, price])
        self.assertEqual(len(unique), 1)
        seen: set[str] = set()
        self.assertFalse(is_duplicate_product(dummy, seen))
        self.assertTrue(is_duplicate_product(price, seen))

    def test_openprices_keeps_separate_shelf_observations(self) -> None:
        from twextract.ecommerce import drop_duplicate_products, product_from_openprices, product_identity_keys

        first = product_from_openprices({**OPEN_PRICE_ITEM, "id": 101})
        second = product_from_openprices({**OPEN_PRICE_ITEM, "id": 202, "price": 3.99, "date": "2026-09-04"})
        assert first is not None and second is not None
        self.assertIn("obs:101", product_identity_keys(first))
        self.assertFalse(product_identity_keys(first) & product_identity_keys(second))
        self.assertEqual(len(drop_duplicate_products([first, second])), 2)


class GeoUniqueTests(unittest.TestCase):
    def test_geo_identity_is_exit_ip(self) -> None:
        from twextract.geo_scrape import geo_identity_keys

        first = {"exit": {"ip": "1.2.3.4"}, "quote": {"text": "hello"}}
        second = {"exit": {"ip": "1.2.3.4"}, "quote": {"text": "other"}}
        third = {"exit": {"ip": "5.6.7.8"}}
        self.assertEqual(geo_identity_keys(first), geo_identity_keys(second))
        self.assertNotEqual(geo_identity_keys(first), geo_identity_keys(third))
        self.assertEqual(geo_identity_keys({}), set())

    def test_proxy_407_stops_the_geo_loop(self) -> None:
        from unittest.mock import patch

        from twextract.errors import GatewayError
        from twextract.geo_scrape import run_geo_batch

        auth_error = GatewayError("proxy inactive", status_code=407)
        with patch("twextract.geo_scrape.inspect_exit_ip", side_effect=auth_error) as inspect:
            with self.assertRaises(GatewayError) as caught:
                run_geo_batch(
                    object(),
                    countries=["us"],
                    count=1000,
                    pause_seconds=0,
                    skip_quotes=True,
                    skip_steam=True,
                )
        self.assertEqual(caught.exception.status_code, 407)
        self.assertEqual(inspect.call_count, 1)

    def test_empty_ip_stops_after_consecutive_failures(self) -> None:
        from unittest.mock import patch

        from twextract.errors import ExtractError
        from twextract.geo_scrape import run_geo_batch

        with patch("twextract.geo_scrape.inspect_exit_ip", return_value={"ip": ""}):
            with self.assertRaises(ExtractError) as caught:
                run_geo_batch(
                    object(),
                    countries=["us"],
                    count=10,
                    pause_seconds=0,
                    skip_quotes=True,
                    skip_steam=True,
                    max_consecutive_failures=3,
                )
        self.assertIn("failed IP checks", str(caught.exception))

    def test_duplicate_retries_cycle_country_and_new_session(self) -> None:
        from unittest.mock import patch

        from twextract.geo_scrape import run_geo_batch

        countries: list[str] = []
        sessions: list[str] = []

        def fake_inspect(_client, targeting):
            normalized = targeting.normalized()
            countries.append(normalized.country or "")
            sessions.append(normalized.session or "")
            if len(countries) <= 3:
                return {"ip": f"10.0.0.{len(countries)}"}
            return {"ip": "10.0.0.1"}

        with patch("twextract.geo_scrape.inspect_exit_ip", side_effect=fake_inspect):
            rows = run_geo_batch(
                object(),
                countries=["us", "de", "in"],
                count=4,
                pause_seconds=0,
                skip_quotes=True,
                skip_steam=True,
                max_attempts=6,
            )
        self.assertEqual(len(rows), 3)
        self.assertEqual(countries[:6], ["us", "de", "in", "us", "de", "in"])
        self.assertEqual(len(set(sessions)), 6)
        self.assertTrue(all(sessions))


class EcommerceExcelTests(unittest.TestCase):
    def test_write_workbook(self) -> None:
        rows = [
            {
                "row_number": 1,
                "ok": True,
                "use_case": "eCommerce & Price Monitoring",
                "source": "dummyjson.com",
                "category": "smartphones",
                "product_name": "iPhone 5s",
                "price": "$199.99",
                "price_value": 199.99,
                "currency": "USD",
                "in_stock": True,
                "stock_count": 25,
                "collected_at": "2026-09-04T17:00:00Z",
                "exit": {"ip": "1.2.3.4", "country": "United States", "city": "Austin"},
            }
        ]
        summary = [{"metric": "Products scraped", "value": 1}]
        flat = flatten_ecommerce_rows(rows)
        self.assertEqual(flat[0]["product_name"], "iPhone 5s")
        self.assertEqual(flat[0]["exit.ip"], "1.2.3.4")
        self.assertEqual(flat[0]["category"], "smartphones")
        with TemporaryDirectory() as tmp:
            folder = Path(tmp)
            excel_export.OUTPUT_DIR = folder
            excel_export.ECOMMERCE_XLSX = folder / "eCommerce.xlsx"
            path = write_ecommerce_excel(rows, summary)
            from openpyxl import load_workbook

            book = load_workbook(path)
            self.assertEqual(path.name, "eCommerce.xlsx")
            self.assertIn("eCommerce", book.sheetnames)
            self.assertIn("Price summary", book.sheetnames)
            headers = [cell.value for cell in book["eCommerce"][1]]
            self.assertIn("product_name", headers)
            self.assertIn("price", headers)
            self.assertIn("category", headers)
            self.assertIn("exit.ip", headers)
            self.assertEqual(headers[-1], "complete Json")


if __name__ == "__main__":
    unittest.main()
