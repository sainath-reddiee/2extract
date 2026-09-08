from __future__ import annotations

import json
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from twextract import excel_export
from twextract.excel_export import flatten_record, ingest_json_file, json_to_table, write_geo_excel


class JsonTableTests(unittest.TestCase):
    SAMPLE = {
        "row_number": 2,
        "collected_at": "2026-09-04T17:51:51.293841+00:00",
        "label": "US",
        "targeting": {"country": "us", "state": None, "city": None, "zip": None},
        "ok": True,
        "exit": {
            "ip": "71.225.21.204",
            "country": "United States",
            "raw": {"isp": "Comcast Cable Communications, LLC", "query": "71.225.21.204"},
        },
        "quote": {
            "text": "It takes a great deal of bravery",
            "author": "J.K. Rowling",
            "tags": ["courage", "friends"],
            "page": 2,
        },
    }

    def test_flatten_uses_dotted_json_keys(self) -> None:
        flat = flatten_record(self.SAMPLE)
        self.assertEqual(flat["row_number"], 2)
        self.assertEqual(flat["targeting.country"], "us")
        self.assertIsNone(flat["targeting.state"])
        self.assertEqual(flat["exit.ip"], "71.225.21.204")
        self.assertEqual(flat["exit.raw.isp"], "Comcast Cable Communications, LLC")
        self.assertEqual(flat["quote.author"], "J.K. Rowling")
        self.assertEqual(flat["quote.tags"], "courage, friends")
        self.assertEqual(flat["ok"], True)

    def test_table_has_one_excel_row_per_json_object(self) -> None:
        records = [self.SAMPLE, {**self.SAMPLE, "row_number": 3, "quote": {"author": "Mother Teresa", "tags": []}}]
        columns, rows = json_to_table(records)
        self.assertEqual(len(rows), 2)
        self.assertIn("quote.author", columns)
        self.assertIn("targeting.zip", columns)
        self.assertEqual(rows[0]["quote.author"], "J.K. Rowling")
        self.assertEqual(rows[1]["quote.author"], "Mother Teresa")
        self.assertEqual(set(rows[0]), set(columns))
        self.assertEqual(set(rows[1]), set(columns))
        self.assertEqual(columns[-1], "complete Json")
        self.assertIn("J.K. Rowling", rows[0]["complete Json"])
        self.assertIn("Mother Teresa", rows[1]["complete Json"])

    def test_upcoming_key_becomes_a_new_column(self) -> None:
        first = {"id": 1, "name": "a"}
        second = {"id": 2, "name": "b", "steam": {"price": "$9.99"}}
        columns, rows = json_to_table([first, second])
        self.assertIn("steam.price", columns)
        self.assertEqual(rows[0]["steam.price"], None)
        self.assertEqual(rows[1]["steam.price"], "$9.99")
        self.assertEqual(len(rows), 2)

    def test_write_headers_are_json_keys(self) -> None:
        with TemporaryDirectory() as tmp:
            folder = Path(tmp)
            excel_export.OUTPUT_DIR = folder
            excel_export.GEO_ADDRESS_XLSX = folder / "geo_address.xlsx"
            path = write_geo_excel([self.SAMPLE, {**self.SAMPLE, "row_number": 3}])
            from openpyxl import load_workbook

            book = load_workbook(path)
            self.assertEqual(path.name, "geo_address.xlsx")
            sheet = book["geo_address"]
            headers = [cell.value for cell in sheet[1]]
            self.assertEqual(sheet.max_row, 3)
            self.assertIn("exit.ip", headers)
            self.assertIn("quote.author", headers)
            self.assertIn("targeting.country", headers)
            self.assertEqual(headers[-1], "complete Json")
            ip_col = headers.index("exit.ip") + 1
            author_col = headers.index("quote.author") + 1
            json_col = headers.index("complete Json") + 1
            self.assertEqual(sheet.cell(2, ip_col).value, "71.225.21.204")
            self.assertEqual(sheet.cell(2, author_col).value, "J.K. Rowling")
            self.assertIn("71.225.21.204", sheet.cell(2, json_col).value)

    def test_ingest_json_file(self) -> None:
        with TemporaryDirectory() as tmp:
            folder = Path(tmp)
            excel_export.OUTPUT_DIR = folder
            excel_export.GEO_ADDRESS_XLSX = folder / "geo_address.xlsx"
            json_path = folder / "geo_results.json"
            json_path.write_text(json.dumps([self.SAMPLE] * 5), encoding="utf-8")
            path = ingest_json_file(json_path, "geo_address")
            from openpyxl import load_workbook

            self.assertEqual(path.name, "geo_address.xlsx")
            sheet = load_workbook(path)["geo_address"]
            self.assertEqual(sheet.max_row, 6)

    def test_strips_excel_illegal_control_characters(self) -> None:
        from openpyxl import load_workbook

        from twextract.excel_export import write_ecommerce_excel

        dirty = "VINAIGRETTE\x01 PROVENALE BALSAMIQUE 36CL"
        rows = [{"product_name": dirty, "ok": True}]
        with TemporaryDirectory() as tmp:
            folder = Path(tmp)
            excel_export.OUTPUT_DIR = folder
            excel_export.ECOMMERCE_XLSX = folder / "eCommerce.xlsx"
            path = write_ecommerce_excel(rows)
            sheet = load_workbook(path)["eCommerce"]
            headers = [cell.value for cell in sheet[1]]
            col = headers.index("product_name") + 1
            self.assertEqual(sheet.cell(2, col).value, "VINAIGRETTE PROVENALE BALSAMIQUE 36CL")


if __name__ == "__main__":
    unittest.main()
