"""Convert scrape JSON into a 1-row-per-record Excel table.

Every JSON key becomes a column. Nested objects use dotted paths
(exit.ip, targeting.country, quote.tags). New keys on later records
are appended as extra columns so upcoming fields are not dropped.
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from twextract.config import ROOT

OUTPUT_DIR = ROOT / "data" / "output"
GEO_ADDRESS_XLSX = OUTPUT_DIR / "geo_address.xlsx"
ECOMMERCE_XLSX = OUTPUT_DIR / "eCommerce.xlsx"
MARKET_RESEARCH_XLSX = OUTPUT_DIR / "market_research.xlsx"
MOBILE_XLSX = OUTPUT_DIR / "mobile.xlsx"
CONNECTIVITY_XLSX = OUTPUT_DIR / "connectivity.xlsx"
# Kept for older tests/scripts that patched this name.
MASTER_XLSX = GEO_ADDRESS_XLSX

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
CELL_FONT = Font(name="Calibri", size=11)
OK_FILL = PatternFill("solid", fgColor="C6EFCE")
FAIL_FILL = PatternFill("solid", fgColor="FFC7CE")
THIN = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9"),
)
WRAP = Alignment(vertical="center", wrap_text=True)
LEFT = Alignment(vertical="center", horizontal="left")

SUMMARY_COLUMNS = [
    ("metric", "metric"),
    ("value", "value"),
]
COMPLETE_JSON_COLUMN = "complete Json"


def flatten_record(value: Any, prefix: str = "") -> dict[str, Any]:
    """Turn one JSON object into {dotted.key: value} columns."""
    items: dict[str, Any] = {}
    if isinstance(value, dict):
        for key, child in value.items():
            name = f"{prefix}.{key}" if prefix else str(key)
            if isinstance(child, dict):
                items.update(flatten_record(child, name))
            elif isinstance(child, list):
                items[name] = _list_cell(child)
            else:
                items[name] = child
        return items
    if prefix:
        items[prefix] = value
    return items


def extract_records(payload: Any) -> list[dict[str, Any]]:
    """Accept a list, or a wrapper object like {products: [...], summary: [...]}."""
    if isinstance(payload, list):
        return [item for item in payload if isinstance(item, dict)]
    if isinstance(payload, dict):
        for key in ("products", "rows", "data", "items", "results"):
            nested = payload.get(key)
            if isinstance(nested, list) and nested and isinstance(nested[0], dict):
                return nested
        return [payload]
    return []


def json_to_table(records: list[dict[str, Any]]) -> tuple[list[str], list[dict[str, Any]]]:
    """Build a rectangular table: union of keys, one aligned row per JSON object."""
    flats = [flatten_record(record) for record in records]
    columns: list[str] = []
    seen: set[str] = set()
    for flat in flats:
        for key in flat:
            if key not in seen and key != COMPLETE_JSON_COLUMN:
                seen.add(key)
                columns.append(key)
    columns.append(COMPLETE_JSON_COLUMN)
    aligned: list[dict[str, Any]] = []
    for record, flat in zip(records, flats, strict=True):
        row = {key: flat.get(key) for key in columns if key != COMPLETE_JSON_COLUMN}
        row[COMPLETE_JSON_COLUMN] = _complete_json(record)
        aligned.append(row)
    return columns, aligned


def write_records_excel(
    records: list[dict[str, Any]],
    sheet_name: str,
    stamp_prefix: str | None = None,
    extra_sheets: dict[str, tuple[list[tuple[str, str]], list[dict[str, Any]]]] | None = None,
    dest: Path | None = None,
) -> Path:
    columns, rows = json_to_table(records)
    sheet_spec = {sheet_name: ([(key, key) for key in columns], rows)}
    if extra_sheets:
        sheet_spec.update(extra_sheets)
    path = dest or _workbook_path_for(sheet_name, stamp_prefix or _safe_name(sheet_name))
    return _write_workbook(path, sheet_spec)


def write_geo_excel(rows: list[dict[str, Any]]) -> Path:
    return write_records_excel(rows, "geo_address", dest=GEO_ADDRESS_XLSX)


def write_ecommerce_excel(rows: list[dict[str, Any]], summary: list[dict[str, Any]] | None = None) -> Path:
    extra = None
    if summary:
        extra = {"Price summary": (SUMMARY_COLUMNS, summary)}
    return write_records_excel(rows, "eCommerce", extra_sheets=extra, dest=ECOMMERCE_XLSX)


def write_market_research_excel(rows: list[dict[str, Any]], summary: list[dict[str, Any]] | None = None) -> Path:
    extra = None
    if summary:
        extra = {"Research summary": (SUMMARY_COLUMNS, summary)}
    return write_records_excel(rows, "Market Research", extra_sheets=extra, dest=MARKET_RESEARCH_XLSX)


def write_mobile_excel(result: dict[str, Any]) -> Path:
    return write_records_excel([result], "Mobile scrape", dest=MOBILE_XLSX)


def write_connectivity_excel(result: dict[str, Any]) -> Path:
    rows = []
    rotating = result.get("rotating")
    if isinstance(rotating, dict):
        rows.append({"mode": "rotating", **rotating, "gateway_protocol": result.get("gateway_protocol")})
    geo = result.get("geo")
    if isinstance(geo, dict):
        rows.append({"mode": "geo", "targeting": geo.get("targeting"), **(geo.get("exit") or {}), "gateway_protocol": result.get("gateway_protocol")})
    if not rows:
        rows = [result]
    return write_records_excel(rows, "Connectivity", dest=CONNECTIVITY_XLSX)


def ingest_json_file(json_path: Path, sheet_name: str | None = None) -> Path:
    payload = json.loads(json_path.read_text(encoding="utf-8"))
    records = extract_records(payload)
    name = sheet_name or _sheet_name_for_file(json_path, payload)
    extra = None
    if isinstance(payload, dict) and isinstance(payload.get("summary"), list):
        if "market" in json_path.name.lower() or "research" in json_path.name.lower():
            extra = {"Research summary": (SUMMARY_COLUMNS, payload["summary"])}
        else:
            extra = {"Price summary": (SUMMARY_COLUMNS, payload["summary"])}
    dest = _workbook_path_for(name, json_path.stem)
    return write_records_excel(records, name, extra_sheets=extra, dest=dest)


def ingest_output_folder(folder: Path | None = None) -> list[Path]:
    folder = folder or OUTPUT_DIR
    mapping = (
        ("geo_results.json", "geo_address"),
        ("ecommerce_results.json", "eCommerce"),
        ("market_research_results.json", "Market Research"),
        ("mobile_results.json", "Mobile scrape"),
    )
    written: list[Path] = []
    for filename, sheet in mapping:
        path = folder / filename
        if path.exists():
            written.append(ingest_json_file(path, sheet))
    return written


def flatten_geo_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    _columns, aligned = json_to_table(rows)
    return aligned


def flatten_ecommerce_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    _columns, aligned = json_to_table(rows)
    return aligned


def flatten_dict(value: Any, prefix: str = "") -> dict[str, Any]:
    return flatten_record(value, prefix)


def _complete_json(record: dict[str, Any]) -> str:
    return _excel_text(json.dumps(record, ensure_ascii=False))


def _list_cell(items: list[Any]) -> Any:
    if not items:
        return ""
    if all(not isinstance(item, (dict, list)) for item in items):
        return _excel_text(", ".join(str(item) for item in items))
    return _excel_text(json.dumps(items, ensure_ascii=False))


def _sheet_name_for_file(path: Path, payload: Any) -> str:
    if path.name.startswith("geo"):
        return "geo_address"
    if path.name.startswith("ecommerce"):
        return "eCommerce"
    if path.name.startswith("market"):
        return "Market Research"
    if path.name.startswith("mobile"):
        return "Mobile scrape"
    if isinstance(payload, dict) and "products" in payload:
        return "eCommerce"
    return path.stem[:31]


def _safe_name(sheet_name: str) -> str:
    return "".join(ch if ch.isalnum() else "_" for ch in sheet_name).strip("_").lower() or "sheet"


def _workbook_path_for(sheet_name: str, stamp_prefix: str) -> Path:
    lowered = f"{sheet_name} {stamp_prefix}".lower()
    if "geo" in lowered:
        return GEO_ADDRESS_XLSX
    if "ecommerce" in lowered or "e-commerce" in lowered:
        return ECOMMERCE_XLSX
    if "market" in lowered or "research" in lowered:
        return MARKET_RESEARCH_XLSX
    if "mobile" in lowered:
        return MOBILE_XLSX
    if "connect" in lowered:
        return CONNECTIVITY_XLSX
    return OUTPUT_DIR / f"{_safe_name(stamp_prefix)}.xlsx"


def _write_workbook(
    path: Path,
    sheets: dict[str, tuple[list[tuple[str, str]], list[dict[str, Any]]]],
) -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    default = workbook.active
    if default is not None:
        workbook.remove(default)
    for sheet_name, (columns, rows) in sheets.items():
        sheet = workbook.create_sheet(sheet_name)
        _render_sheet(sheet, columns, rows)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def _render_sheet(sheet: Worksheet, columns: list[tuple[str, str]], rows: list[dict[str, Any]]) -> None:
    keys = [key for key, _ in columns]
    headers = [title for _, title in columns]
    sheet.append(headers)
    body = rows or [{key: "" for key in keys}]
    for row in body:
        sheet.append([_cell(row.get(key)) for key in keys])

    last_row = sheet.max_row
    last_col = max(sheet.max_column, 1)
    last_letter = get_column_letter(last_col)

    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        cell.border = THIN

    success_cols = {index for index, (key, _) in enumerate(columns, start=1) if key.split(".")[-1] == "ok"}

    for row_idx, row_cells in enumerate(sheet.iter_rows(min_row=2, max_row=last_row, max_col=last_col), start=2):
        for cell in row_cells:
            cell.font = CELL_FONT
            cell.border = THIN
            cell.alignment = WRAP if cell.column in _wrap_columns(columns) else LEFT
        for success_col in success_cols:
            flag = sheet.cell(row_idx, success_col)
            if flag.value in (True, "Yes", "true", "TRUE"):
                flag.fill = OK_FILL
            elif flag.value in (False, "No", "false", "FALSE"):
                flag.fill = FAIL_FILL

    table = Table(displayName=_table_name(sheet.title), ref=f"A1:{last_letter}{last_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(table)
    sheet.auto_filter.ref = f"A1:{last_letter}{last_row}"
    sheet.freeze_panes = "A2"
    sheet.row_dimensions[1].height = 24
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"

    for index, (key, title) in enumerate(columns, start=1):
        width = max(len(str(title)), 10)
        sample_last = min(last_row, 40)
        for row_idx in range(2, sample_last + 1):
            value = sheet.cell(row_idx, index).value
            if value is not None:
                width = max(width, min(len(str(value)), 36))
        cap = 80 if key == COMPLETE_JSON_COLUMN or "text" in key.lower() or "description" in key.lower() else 42
        sheet.column_dimensions[get_column_letter(index)].width = min(width + 2, cap)


def _wrap_columns(columns: list[tuple[str, str]]) -> set[int]:
    markers = ("error", "url", "text", "lead", "tags", "description", "org", "asn", "json", "body", "title")
    wrapped: set[int] = set()
    for index, (key, _title) in enumerate(columns, start=1):
        lowered = key.lower()
        if any(marker in lowered for marker in markers):
            wrapped.add(index)
    return wrapped


def _table_name(sheet_title: str) -> str:
    cleaned = "".join(ch if ch.isalnum() else "_" for ch in sheet_title)
    return f"T_{cleaned}"[:30]


def _excel_text(value: str) -> str:
    """Excel rejects ASCII control characters (except tab/newline/CR) and strings over 32,767 chars."""
    cleaned = ILLEGAL_CHARACTERS_RE.sub("", value)
    limit = 32767
    if len(cleaned) > limit:
        return cleaned[: limit - 18] + "...[truncated]"
    return cleaned


def _cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value
    if not isinstance(value, str):
        value = str(value)
    return _excel_text(value)
